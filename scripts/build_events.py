#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pdfplumber
import requests
from dateutil import parser as date_parser
from icalendar import Calendar
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "events.json"
CONFIG_FILE = ROOT / "config" / "sources.json"
TEAM = "CB BUTchers"
PRAGUE = ZoneInfo("Europe/Prague")


def norm(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip().casefold()


def stable_id(prefix: str, *parts: object) -> str:
    raw = "|".join(str(x) for x in parts)
    return f"{prefix}-{hashlib.sha1(raw.encode('utf-8')).hexdigest()[:12]}"


def base_event(*, event_id, event_type, title, date_str, start_time=None, end_time=None,
               location=None, rink=None, competition=None, description=None, source_type,
               source_name, opponent=None, all_day=False, editable=False):
    return {
        "id": event_id,
        "type": event_type,
        "title": title,
        "team": TEAM,
        "opponent": opponent,
        "date": date_str,
        "startTime": start_time,
        "endTime": end_time,
        "allDay": all_day,
        "location": location,
        "rink": rink,
        "competition": competition,
        "description": description,
        "source": {"type": source_type, "name": source_name, "url": None},
        "editable": editable,
        "attendanceEnabled": True,
    }


def parse_excel(path: Path, cfg: dict) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[cfg.get("sheet", "Rozpis")]
    current_date = None
    out = []

    def parse_header_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            m = re.search(r"(\d{1,2})\.\s*(\d{1,2})\.?\s*(\d{4})", value)
            if m:
                return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        return None

    row = 1
    while row <= ws.max_row:
        col_a = ws.cell(row, 1).value
        col_b = ws.cell(row, 2).value
        if col_a == "Dráha - Čas":
            current_date = parse_header_date(col_b)
            row += 1
            continue

        if isinstance(col_a, str) and re.match(r"^[AB]\s*-\s*\d{1,2}:\d{2}$", col_a):
            team1 = str(col_b or "").strip()
            team2 = str(ws.cell(row + 1, 2).value or "").strip()
            if current_date and (norm(team1) == norm(TEAM) or norm(team2) == norm(TEAM)):
                rink, start = [x.strip() for x in col_a.split("-", 1)]
                opponent = team2 if norm(team1) == norm(TEAM) else team1
                d = current_date.isoformat()
                out.append(base_event(
                    event_id=stable_id("excel", d, start, rink, opponent),
                    event_type="match",
                    title=f"{TEAM} vs {opponent}",
                    date_str=d,
                    start_time=start,
                    location=cfg.get("location", "Curling Brno"),
                    rink=rink,
                    competition=cfg.get("competition", "Brněnský pohár"),
                    source_type="excel",
                    source_name="Brněnský pohár 2026/27",
                    opponent=opponent,
                ))
            row += 2
            continue
        row += 1
    return out


def _split_pdf_columns(words: list[dict], anchors: list[float]) -> dict[int, str]:
    bounds = [(anchors[i] + anchors[i + 1]) / 2 for i in range(len(anchors) - 1)]
    cols = {i + 1: [] for i in range(len(anchors))}
    for w in sorted(words, key=lambda x: x["x0"]):
        idx = 1
        while idx <= len(bounds) and w["x0"] >= bounds[idx - 1]:
            idx += 1
        cols[idx].append(w["text"])
    return {k: " ".join(v).strip() for k, v in cols.items()}


def parse_pdf(path: Path, cfg: dict) -> list[dict]:
    out = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            lines = page.extract_text_lines()
            words = page.extract_words()
            date_headers = []
            for line in lines:
                text = line["text"]
                if not text.startswith("CURLINGOVÁ HALA"):
                    continue
                m = re.search(
                    r"CURLINGOVÁ HALA\s+(.+?),\s+(?:SOBOTA|NEDĚLE)\s+(\d{1,2})\.\s*(\d{1,2})\.?\s*(\d{4})",
                    text, re.I,
                )
                if m:
                    location = "Curlingová hala " + m.group(1).strip().title()
                    d = date(int(m.group(4)), int(m.group(3)), int(m.group(2))).isoformat()
                    date_headers.append((line["top"], d, location))
            if not date_headers:
                continue

            header_lines = [line for line in lines if "1. dráha" in line["text"]]
            for line in lines:
                mt = re.search(r"(\d{1,2}:\d{2})\s*/\s*(\d{1,2}:\d{2})", line["text"])
                if not mt:
                    continue
                dheader = max((h for h in date_headers if h[0] < line["top"]), default=None, key=lambda x: x[0])
                if not dheader:
                    continue
                date_top, d, location = dheader
                hline = min((h for h in header_lines if date_top < h["top"] < line["top"]), default=None, key=lambda x: x["top"])
                if not hline:
                    continue
                hwords = [w for w in words if abs(w["top"] - hline["top"]) < 1.2 and w["x0"] > 100]
                anchors = sorted((int(w["text"][0]), w["x0"]) for w in hwords if re.fullmatch(r"[1-4]\.", w["text"]))
                xs = [x for _, x in anchors]
                if not xs:
                    continue

                t = line["top"]
                row_words = [w for w in words if t - 8.5 <= w["top"] <= t + 15.0 and w["x0"] > 100]
                y_values = sorted(set(round(w["top"], 1) for w in row_words))
                groups = [(y, [w for w in row_words if round(w["top"], 1) == y]) for y in y_values]
                above = [g for g in groups if g[0] < t - 1]
                below = [g for g in groups if g[0] > t + 1]
                if not above or not below:
                    continue
                top_cols = _split_pdf_columns(max(above, key=lambda g: g[0])[1], xs)
                bottom_cols = _split_pdf_columns(min(below, key=lambda g: g[0])[1], xs)
                for rink in range(1, len(xs) + 1):
                    team1 = top_cols.get(rink, "")
                    team2 = bottom_cols.get(rink, "")
                    if norm(team1) != norm(TEAM) and norm(team2) != norm(TEAM):
                        continue
                    opponent = team2 if norm(team1) == norm(TEAM) else team1
                    start = mt.group(2)
                    out.append(base_event(
                        event_id=stable_id("pdf", d, start, rink, opponent),
                        event_type="match",
                        title=f"{TEAM} vs {opponent}",
                        date_str=d,
                        start_time=start,
                        location=location,
                        rink=str(rink),
                        competition=cfg.get("competition", "MČR divize"),
                        description=f"Zdroj: PDF, strana {page_no}.",
                        source_type="pdf",
                        source_name="MČR divize 2026/27",
                        opponent=opponent,
                    ))
    return out


def _dt_parts(value):
    if isinstance(value, datetime):
        dt = value if value.tzinfo else value.replace(tzinfo=PRAGUE)
        dt = dt.astimezone(PRAGUE)
        return dt.date().isoformat(), dt.strftime("%H:%M"), False
    if isinstance(value, date):
        return value.isoformat(), None, True
    dt = date_parser.parse(str(value))
    if dt.tzinfo:
        dt = dt.astimezone(PRAGUE)
    return dt.date().isoformat(), dt.strftime("%H:%M"), False


def parse_ical(url: str, cfg: dict) -> list[dict]:
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    cal = Calendar.from_ical(response.content)
    out = []
    for component in cal.walk():
        if component.name != "VEVENT":
            continue
        summary = str(component.get("summary", "Událost"))
        location = str(component.get("location", "")).strip() or None
        description = str(component.get("description", "")).strip() or None
        uid = str(component.get("uid", "")) or stable_id("uid", summary, component.get("dtstart"))
        dtstart = component.decoded("dtstart")
        d, start, all_day = _dt_parts(dtstart)
        end = None
        if component.get("dtend"):
            _, end, _ = _dt_parts(component.decoded("dtend"))
        event_type = "training" if re.search(r"tr[eé]nink|training", summary, re.I) else "event"
        out.append(base_event(
            event_id=stable_id("ical", uid),
            event_type=event_type,
            title=summary,
            date_str=d,
            start_time=start,
            end_time=end,
            location=location,
            competition=None,
            description=description,
            source_type="ical",
            source_name=cfg.get("name", "Curling Brno EOS"),
            all_day=all_day,
        ))
    return out


def load_existing() -> list[dict]:
    if not DATA_FILE.exists():
        return []
    return json.loads(DATA_FILE.read_text(encoding="utf-8"))


def source_slice(existing, source_type):
    return [e for e in existing if e.get("source", {}).get("type") == source_type]


def maybe_download_excel(local_path: Path) -> Path:
    url = os.getenv("GOOGLE_SHEET_XLSX_URL", "").strip()
    if not url:
        return local_path
    try:
        r = requests.get(url, timeout=30)
        r.raise_for_status()
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.write(r.content)
        tmp.close()
        print("Excel stažen z GOOGLE_SHEET_XLSX_URL")
        return Path(tmp.name)
    except Exception as exc:
        print(f"VAROVÁNÍ: Google Sheet se nepodařilo stáhnout, používám lokální Excel: {exc}")
        return local_path


def main():
    cfg = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    existing = load_existing()

    try:
        excel_path = maybe_download_excel(ROOT / cfg["excel"]["path"])
        excel_events = parse_excel(excel_path, cfg["excel"])
        print(f"Excel: {len(excel_events)} zápasů")
    except Exception as exc:
        print(f"VAROVÁNÍ: Excel import selhal, zachovávám poslední data: {exc}")
        excel_events = source_slice(existing, "excel")

    try:
        pdf_events = parse_pdf(ROOT / cfg["pdf"]["path"], cfg["pdf"])
        print(f"PDF: {len(pdf_events)} zápasů")
    except Exception as exc:
        print(f"VAROVÁNÍ: PDF import selhal, zachovávám poslední data: {exc}")
        pdf_events = source_slice(existing, "pdf")

    ical_url = os.getenv(cfg.get("ical", {}).get("urlFromEnvironment", "ICAL_URL"), "").strip()
    if ical_url:
        try:
            ical_events = parse_ical(ical_url, cfg.get("ical", {}))
            print(f"iCal: {len(ical_events)} událostí")
        except Exception as exc:
            print(f"VAROVÁNÍ: iCal import selhal, zachovávám poslední data: {exc}")
            ical_events = source_slice(existing, "ical")
    else:
        print("iCal: proměnná ICAL_URL není nastavena, zachovávám poslední data.")
        ical_events = source_slice(existing, "ical")

    merged = excel_events + pdf_events + ical_events
    dedup = {e["id"]: e for e in merged}
    events = sorted(dedup.values(), key=lambda e: (e.get("date", ""), e.get("startTime") or "", e.get("title", "")))
    DATA_FILE.write_text(json.dumps(events, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Celkem: {len(events)} událostí -> {DATA_FILE.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
